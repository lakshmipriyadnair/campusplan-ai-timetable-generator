"""Strict .xlsx templates and preview/apply imports for source academic data."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app import models, schemas
from app.database import get_db
from app.predefined_groups import (
    PredefinedGroupConfigurationError,
    PredefinedGroupDefinition,
    configure_predefined_groups,
)


router = APIRouter(prefix="/imports", tags=["Excel imports"])
MAX_FILE_BYTES = 5 * 1024 * 1024

# These columns intentionally mirror existing Pydantic input schemas.
SPECS: dict[str, tuple[list[str], type[Any], type[Any]]] = {
    "faculty": (["faculty_name", "faculty_code", "email"], models.Faculty, schemas.FacultyCreate),
    "rooms": (["room_name", "room_type", "capacity"], models.Room, schemas.RoomCreate),
    "subjects": (["subject_code", "subject_name"], models.Subject, schemas.SubjectCreate),
    "semester-subjects": (
        ["semester_id", "student_group_id", "subject_id", "weekly_lecture_hours",
         "weekly_tutorial_hours", "weekly_practical_hours", "lecture_session_minutes",
         "tutorial_session_minutes", "practical_session_minutes"],
        models.SemesterSubject,
        schemas.SemesterSubjectCreate,
    ),
}
GROUP_COLUMNS = ["group_code", "semester_number", "academic_year_name", "group_name", "student_count"]
UNIQUE_COLUMNS = {
    "faculty": ["faculty_code", "email"],
    "rooms": ["room_name"],
    "subjects": ["subject_code"],
}


def parse_workbook(filename: str, content: bytes, expected_kind: str) -> tuple[str, list[dict[str, Any]]]:
    """Check file type, size, workbook structure, and exact expected columns."""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Choose an Excel .xlsx file.")
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="The .xlsx file must be between 1 byte and 5 MB.")
    if expected_kind not in (*SPECS, "student-groups"):
        raise HTTPException(status_code=404, detail="Unknown import template.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        header_row = next(values, None)
    except Exception as error:
        raise HTTPException(status_code=422, detail="The file is not a readable .xlsx workbook.") from error
    expected = GROUP_COLUMNS if expected_kind == "student-groups" else SPECS[expected_kind][0]
    headers = [str(value).strip() if value is not None else "" for value in (header_row or ())]
    while headers and headers[-1] == "":
        headers.pop()
    if headers != expected:
        raise HTTPException(
            status_code=422,
            detail={"message": "Spreadsheet columns do not match the selected template.",
                    "expected_columns": expected, "received_columns": headers},
        )
    rows: list[dict[str, Any]] = []
    for row_number, raw in enumerate(values, start=2):
        cells = list(raw[:len(expected)])
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        rows.append({column: (cells[index] if index < len(cells) else None)
                     for index, column in enumerate(expected)} | {"_row": row_number})
    workbook.close()
    return sha256(content).hexdigest(), rows


def validate_rows(kind: str, rows: list[dict[str, Any]], db: Session) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: dict[str, set[Any]] = {}

    for row in rows:
        row_number = row["_row"]
        data = {key: value for key, value in row.items() if key != "_row"}
        if kind == "student-groups":
            keys = {
                "group_code": data.get("group_code"),
                "semester/group": (data.get("academic_year_name"), data.get("semester_number"), data.get("group_name")),
            }
        elif kind in UNIQUE_COLUMNS:
            keys = {key: data.get(key) for key in UNIQUE_COLUMNS[kind] if data.get(key) not in (None, "")}
        elif kind == "semester-subjects":
            keys = {"group/subject": (data.get("student_group_id"), data.get("subject_id"))}
        else:
            keys = {}
        duplicate_found = False
        for key, value in keys.items():
            if value in seen.setdefault(key, set()):
                errors.append({"row": row_number, "column": key, "message": f"Duplicate {key} in this file."})
                duplicate_found = True
            seen[key].add(value)
            if kind in UNIQUE_COLUMNS and db.query(SPECS[kind][1]).filter(getattr(SPECS[kind][1], key) == value).first():
                errors.append({"row": row_number, "column": key, "message": f"{key} already exists; imports never overwrite records."})
                duplicate_found = True
        try:
            if kind == "student-groups":
                item = PredefinedGroupDefinition.model_validate(data)
                data = item.model_dump()
            else:
                schema = SPECS[kind][2]
                if kind == "faculty" and data.get("email") in (None, ""):
                    data["email"] = None
                if kind == "semester-subjects":
                    for key in ("lecture_session_minutes", "tutorial_session_minutes", "practical_session_minutes"):
                        if data.get(key) in (None, ""):
                            data[key] = None
                item = schema.model_validate(data)
                data = item.model_dump()
        except ValidationError as error:
            for issue in error.errors():
                field = ".".join(str(part) for part in issue["loc"])
                errors.append({"row": row_number, "column": field,
                               "message": issue["msg"].replace("Value error, ", "")})
            continue

        if duplicate_found:
            continue

        if kind == "student-groups":
            semester = db.scalar(
                select(models.Semester)
                .join(models.Semester.department)
                .join(models.Semester.academic_year)
                .where(
                    models.Department.code == "CSE",
                    models.AcademicYear.name == data["academic_year_name"],
                    models.Semester.semester_number == data["semester_number"],
                )
            )
            if semester is None:
                errors.append({"row": row_number, "column": "semester_number",
                               "message": "No matching CSE semester and academic year exists."})
                continue
            identity = db.scalar(select(models.PredefinedStudentGroup).where(
                models.PredefinedStudentGroup.group_code == data["group_code"]
            ))
            if identity is not None:
                existing_group = db.get(models.StudentGroup, identity.student_group_id)
                if existing_group is None or (
                    existing_group.semester_id != semester.id
                    or existing_group.group_name != data["group_name"]
                    or existing_group.student_count != data["student_count"]
                ):
                    errors.append({"row": row_number, "column": "group_code",
                                   "message": "This code already identifies different protected group data."})
                    continue
            else:
                existing_group = db.scalar(select(models.StudentGroup).where(
                    models.StudentGroup.semester_id == semester.id,
                    models.StudentGroup.group_name == data["group_name"],
                ))
                if existing_group is not None and existing_group.student_count != data["student_count"]:
                    errors.append({"row": row_number, "column": "student_count",
                                   "message": "A group with this name already exists with a different count."})
                    continue

        if kind == "semester-subjects":
            semester = db.get(models.Semester, data["semester_id"])
            group = db.get(models.StudentGroup, data["student_group_id"])
            subject = db.get(models.Subject, data["subject_id"])
            for column, value in (("semester_id", semester), ("student_group_id", group), ("subject_id", subject)):
                if value is None:
                    errors.append({"row": row_number, "column": column, "message": "Referenced record does not exist."})
            if semester and group and semester.id != group.semester_id:
                errors.append({"row": row_number, "column": "student_group_id", "message": "Student group does not belong to this semester."})
            if semester is None or group is None or subject is None or (group and semester and group.semester_id != semester.id):
                continue
            department = db.get(models.Department, semester.department_id) if semester else None
            if (
                semester.semester_number == 5
                and department is not None
                and department.code == "CSE"
                and db.scalar(select(models.SharedSubjectRequirement.id).where(
                    models.SharedSubjectRequirement.semester_id == semester.id,
                    models.SharedSubjectRequirement.subject_id == subject.id,
                ))
            ):
                errors.append({
                    "row": row_number,
                    "column": "subject_id",
                    "message": "A shared S5 requirement already exists; use section-specific teacher assignments instead of importing duplicate weekly hours.",
                })
                continue
            if db.scalar(select(models.SemesterSubject.id).where(
                models.SemesterSubject.student_group_id == group.id,
                models.SemesterSubject.subject_id == subject.id,
            )):
                errors.append({"row": row_number, "column": "subject_id", "message": "This subject is already assigned to the student group."})
                continue
        valid.append(data)
    return valid, errors


@router.get("/templates/{kind}")
def download_template(kind: str) -> StreamingResponse:
    if kind not in (*SPECS, "student-groups"):
        raise HTTPException(status_code=404, detail="Unknown import template.")
    columns = GROUP_COLUMNS if kind == "student-groups" else SPECS[kind][0]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(columns)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{kind}-template.xlsx"'})


@router.post("/{kind}/preview")
async def preview_import(kind: str, file: UploadFile = File(...), db: Session = Depends(get_db)) -> dict[str, Any]:
    content = await file.read(MAX_FILE_BYTES + 1)
    digest, rows = parse_workbook(file.filename or "", content, kind)
    valid, errors = validate_rows(kind, rows, db)
    return {"preview_id": digest, "kind": kind, "row_count": len(rows),
            "valid_count": len(valid), "errors": errors, "rows": valid[:100],
            "can_apply": bool(valid) and not errors, "preview_limit": 100}


@router.post("/{kind}/apply")
async def apply_import(kind: str, expected_preview_id: str = Form(...), file: UploadFile = File(...),
                       db: Session = Depends(get_db)) -> dict[str, Any]:
    content = await file.read(MAX_FILE_BYTES + 1)
    digest, rows = parse_workbook(file.filename or "", content, kind)
    if digest != expected_preview_id:
        raise HTTPException(status_code=409, detail="The file changed after preview. Preview it again before applying.")
    valid, errors = validate_rows(kind, rows, db)
    if errors or not valid:
        raise HTTPException(status_code=422, detail={"message": "Import validation failed; no rows were saved.", "errors": errors})
    try:
        if kind == "student-groups":
            configured = configure_predefined_groups(db, [PredefinedGroupDefinition.model_validate(row) for row in valid])
            return {"imported_count": len(configured), "kind": kind, "message": "Protected groups configured."}
        model = SPECS[kind][1]
        db.add_all([model(**row) for row in valid])
        db.commit()
    except PredefinedGroupConfigurationError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="A record conflicts with data saved since preview. No rows were saved; preview again.") from error
    return {"imported_count": len(valid), "kind": kind, "message": "All validated rows were imported."}
